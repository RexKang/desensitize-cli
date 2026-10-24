#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""对 tests/fixtures 跑完脱敏后的断言检查。

用法：
    python tests/make_demo.py
    python desensitize.py tests/fixtures --yes --rules tests/rules.test.json
    python tests/verify.py
"""

import sys
from pathlib import Path

from openpyxl import load_workbook

# Windows 控制台默认 cp1252/cp936，直接 print 中文会抛 UnicodeEncodeError
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

FIX = Path(__file__).resolve().parent / "fixtures"
sys.path.insert(0, str(FIX.parent.parent))
from desensitize import Rule, detect_encoding  # noqa: E402

ok = True


def chk(label: str, cond, extra: str = "") -> None:
    global ok
    ok = ok and bool(cond)
    print(("PASS  " if cond else "FAIL  ") + label + (f"   [{extra}]" if extra else ""))


# ── 文本：各类规则命中与不命中
txt = (FIX / "sample.txt").read_bytes().decode("utf-8")
chk("字面删除 TCK-", "工单 2026-0001" in txt)
chk("手机号等长星号", "联系人 ***********、user@example.com" in txt)
chk("身份证 18 位星号", "证件 ******************" in txt)
chk("IPv4 第一段固定 + 其余段位移", "内网 203.8.9.10 和 1.10.0.1" in txt)
chk("网段地址同样位移", "网段 203.8.9.7/24" in txt)
chk("带端口保留后缀", "带端口 203.8.9.10:8080" in txt)
chk("带编号保留后缀", "带编号 203.8.9.10-188" in txt)
chk("三段数字不当 IP", "非IP 10.1.2-188" in txt)
chk("192.168 段位移", "私有段 203.175.8.17" in txt)
chk("日期不动", "日期 20260101" in txt)
chk("8 位参数不动", "数值 10000000" in txt)
chk("20 位长串不拆", "长串 12345678901234567890" in txt)
chk("13 位数字不动", "13位 1380013800000" in txt)
chk("前缀替换（词边界命中）", "标签 SITE-SRV01 与 MAP-01" in txt)
chk("主机名已哈希、原值消失", "host-web01" not in txt and "218c248501" in txt)

# ── 编码：GB18030 与 UTF-16 都要原样写回
md_raw = (FIX / "sample-gbk.md").read_bytes()
try:
    md_raw.decode("utf-8")
    still_gbk = False
except UnicodeDecodeError:
    still_gbk = True
chk("GB18030 编码保持", still_gbk)
md = md_raw.decode("gb18030")
chk("md 表格已脱敏", "| 0001 | 甲 | *********** | user@example.com | 203.8.9.10 |" in md)
chk("md 表头未被破坏", "| 单号 | 联系人 | 手机 | 邮箱 | 地址 |" in md)

u16_raw = (FIX / "sample-utf16.txt").read_bytes()
chk("UTF-16 BOM 保留", u16_raw.startswith(b"\xff\xfe"))
u16 = u16_raw[2:].decode("utf-16-le")
chk("UTF-16 内容已脱敏", "***********" in u16 and "13800138000" not in u16)

# ── xlsx：数值格、日期格、公式、复合替换
ws = load_workbook(FIX / "sample.xlsx")["工单"]
chk("xlsx 数值格被替换成星号串", ws["C2"].value == "***********", repr(ws["C2"].value))
chk("xlsx 邮箱整格替换", ws["D2"].value == "user@example.com", ws["D2"].value)
chk("xlsx IP 格位移", ws["E2"].value == "203.8.9.10", ws["E2"].value)
chk("xlsx 192.168 格位移", ws["E3"].value == "203.175.8.17", ws["E3"].value)
chk("xlsx 日期格不动", str(ws["F2"].value)[:10] == "2026-01-01", str(ws["F2"].value))
chk("xlsx 公式保留", ws["H2"].value == "=C2", ws["H2"].value)
chk("xlsx 文本格内嵌替换", ws["G2"].value == "联系 *********** 或 SITE-SRV01", ws["G2"].value)
chk("xlsx 同格多处替换", ws["I2"].value == "1-SITE-2", ws["I2"].value)
chk("xlsx 整格删空为真空格", ws["I3"].value == "ONLY", ws["I3"].value)

# ── 目录处理：递归 + 剪枝
deep = (FIX / "nested" / "l2" / "l3" / "deep.txt").read_bytes().decode("utf-8")
chk("三层嵌套默认被处理", "13800138000" not in deep and "user@example.com" in deep, deep.strip())
hidden = (FIX / ".hidden" / "skip.txt").read_bytes().decode("utf-8")
chk("隐藏目录被跳过", "13800138000" in hidden)

# ── 单元断言：ip_shift / hash / 编码探测
ipr = Rule(0, {"id": "ip", "type": "ip_shift", "first_octet": 203, "offset": 7,
               "source_cidrs": ["10.0.0.0/8"], "replace": ""})
chk("ip_shift 常规位移", ipr.shift_ip("10.1.2.3") == "203.8.9.10")
chk("ip_shift 超 255 取模", ipr.shift_ip("255.255.255.255") == "203.6.6.6")
chk("ip_shift 非法段不处理", ipr.shift_ip("999.1.1.1") == "999.1.1.1")
chk("ip_shift 源网段过滤", ipr.in_source("10.1.2.3") and not ipr.in_source("8.8.8.8"))
ipr2 = Rule(0, {"id": "ip2", "type": "ip_shift", "first_octet": 203,
                "offsets": [None, 1, 2, 3], "replace": ""})
chk("ip_shift 逐段偏移", ipr2.shift_ip("10.1.2.3") == "203.2.4.6")

hr = Rule(0, {"id": "h", "type": "hash", "pattern": "host-[a-z]{2,6}\\d{2,4}",
              "keep_tail": 2, "salt": ""})
h1 = hr.hash_token("host-web01")
chk("hash 长度不变", len(h1) == len("host-web01"), h1)
chk("hash 尾部保留", h1.endswith("01"), h1)
chk("hash 可复现", h1 == hr.hash_token("host-web01"))
chk("hash 不同输入不同结果", hr.hash_token("host-db02") != h1)
hr_salt = Rule(0, {"id": "h", "type": "hash", "pattern": "host-[a-z]{2,6}\\d{2,4}",
                   "keep_tail": 2, "salt": "projectA"})
chk("hash 加盐后映射不同", hr_salt.hash_token("host-web01") != h1)

chk("编码探测 utf-8 BOM", detect_encoding(b"\xef\xbb\xbfabc")[0] == "utf-8")
chk("编码探测 utf-16-le", detect_encoding(b"\xff\xfea\x00")[0] == "utf-16-le")
chk("编码探测 utf-16-be", detect_encoding(b"\xfe\xff\x00a")[0] == "utf-16-be")
chk("编码探测纯 ASCII", detect_encoding(b"plain text")[0] == "utf-8")

print()
print("全部通过" if ok else "有失败项")
sys.exit(0 if ok else 1)
