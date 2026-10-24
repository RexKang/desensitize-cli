#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成自测夹具：txt / GB18030-md / xlsx / 三层嵌套 / 隐藏目录。

夹具里全部是编造的示例数据（example.com、10.0.0.0/8 私网段、顺序号），
不包含任何真实业务信息。每次跑自测前执行一次，保证夹具回到「未脱敏」状态。
"""

from datetime import date
from pathlib import Path

from openpyxl import Workbook

import sys

# Windows 控制台默认 cp1252/cp936，打印路径里的中文会报错
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

FIX = Path(__file__).resolve().parent / "fixtures"

TXT = """工单 TCK-2026-0001
联系人 13800138000、user@example.com
证件 110101199001011234
内网 10.1.2.3 和 1.10.0.1
网段 10.1.2.0/24
带端口 10.1.2.3:8080
带编号 10.1.2.3-188
非IP 10.1.2-188
私有段 192.168.1.10
日期 20260101
数值 10000000
长串 12345678901234567890
13位 1380013800000
标签 AP-SRV01 与 MAP-01
主机 host-web01 与 host-db02
"""

MD = """# 示例数据（GB18030 编码）

| 单号 | 联系人 | 手机 | 邮箱 | 地址 |
|---|---|---|---|---|
| TCK-0001 | 甲 | 13800138000 | user@example.com | 10.1.2.3 |
| TCK-0002 | 乙 | 13900139000 | staff@example.org | 10.9.9.9 |

备注：日期 20260101 保留。
"""


def main() -> None:
    FIX.mkdir(parents=True, exist_ok=True)
    (FIX / "sample.txt").write_bytes(TXT.encode("utf-8"))
    (FIX / "sample-gbk.md").write_bytes(MD.encode("gb18030"))
    # UTF-16 LE + BOM：验证编码探测与写回
    (FIX / "sample-utf16.txt").write_bytes(
        "\ufeff联系人 13800138000、utf16@example.com\n".encode("utf-16-le"))

    # 三层嵌套，验证默认递归
    deep = FIX / "nested" / "l2" / "l3"
    deep.mkdir(parents=True, exist_ok=True)
    (deep / "deep.txt").write_bytes("深层 13800138000、deep@example.com\n".encode("utf-8"))

    # 隐藏目录，验证会被剪掉
    hidden = FIX / ".hidden"
    hidden.mkdir(exist_ok=True)
    (hidden / "skip.txt").write_bytes("不该被处理 13800138000\n".encode("utf-8"))

    wb = Workbook()
    ws = wb.active
    ws.title = "工单"
    ws.append(["单号", "联系人", "手机", "邮箱", "内网IP", "日期格", "文本格", "公式"])
    ws.append(["TCK-0001", "甲", 13800138000, "user@example.com", "10.1.2.3",
               date(2026, 1, 1), "联系 13800138000 或 AP-SRV01", "=C2"])
    ws.append(["TCK-0002", "乙", 13900139000, "staff@example.org", "192.168.1.10",
               date(2026, 1, 2), "13位 1380013800000，日期 20260101", "=C3"])
    ws["I1"] = "复合"
    ws["I2"] = "TCK-1-AP-2"
    ws["I3"] = "TCK-ONLY"
    wb.save(FIX / "sample.xlsx")

    for p in sorted(FIX.rglob("*")):
        if p.is_file():
            print(f"{p.relative_to(FIX)!s:<28} {p.stat().st_size} bytes")


if __name__ == "__main__":
    main()
